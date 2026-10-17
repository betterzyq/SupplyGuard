"""Bounded CSV/XLSX parsing and formula-safe downloadable files."""

import csv
import io
import zipfile
from pathlib import Path
from pydantic import ValidationError
from openpyxl import Workbook, load_workbook
from .schemas import ROWS, references

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000


class SourceRow(dict):
    """Source row location survives ignored blanks without polluting mapped data."""

    def __init__(self, values, number):
        super().__init__(values)
        self.number = number


def read_table(content, filename):
    try:
        return _read_table(content, filename)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Malformed CSV/XLSX file") from exc


def _read_table(content, filename):
    if len(content) > MAX_BYTES:
        raise ValueError("File exceeds 2 MiB")
    data = []

    def append(row, number):
        if number > 10000 or len(row) > 40:
            raise ValueError("Maximum 10000 physical rows / 40 columns")
        if not any(v is not None and str(v).strip() for v in row):
            return
        if len(data) >= MAX_ROWS + 1:
            raise ValueError("Maximum 1000 nonempty data rows")
        if any(len(str(v or "")) > 2000 for v in row):
            raise ValueError(f"Row {number}: oversized value")
        data.append((number, row))

    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        for number, row in enumerate(
            csv.reader(io.StringIO(content.decode("utf-8-sig"))), 1
        ):
            append(row, number)
    elif suffix == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 100
                or sum(f.file_size for f in archive.infolist()) > 10 * 1024 * 1024
            ):
                raise ValueError("Expanded workbook exceeds safety limit")
        book = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(book.worksheets) != 1:
                raise ValueError("Import requires exactly one sheet")
            sheet = book.worksheets[0]
            sheet.reset_dimensions()
            for number, cells in enumerate(sheet.iter_rows(), 1):
                if any(c.data_type == "f" for c in cells):
                    raise ValueError(
                        "Formula cells are not allowed; upload literal values"
                    )
                append(
                    [
                        (
                            c.value.isoformat()
                            if hasattr(c.value, "isoformat")
                            else c.value
                        )
                        for c in cells
                    ],
                    number,
                )
        finally:
            book.close()
    else:
        raise ValueError("Only UTF-8 CSV or .xlsx supported")
    if len(data) < 2:
        raise ValueError("Require header and 1..1000 data rows")
    headers = [str(x or "").strip() for x in data[0][1]]
    if (
        not all(headers)
        or len(set(headers)) != len(headers)
        or any(len(x) > 100 for x in headers)
    ):
        raise ValueError("Headers must be nonempty and unique, up to 100 characters")
    records = []
    for number, row in data[1:]:
        if len(row) > len(headers):
            raise ValueError(f"Row {number}: too many cells")
        records.append(
            SourceRow(
                {
                    h: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
                    for i, h in enumerate(headers)
                },
                number,
            )
        )
    return headers, records


def normalize(kind, headers, records, mapping):
    model = ROWS[kind]
    if (
        not isinstance(mapping, dict)
        or any(
            not isinstance(k, str) or not isinstance(v, str) for k, v in mapping.items()
        )
        or set(mapping) - set(model.model_fields)
        or len(set(mapping.values())) != len(mapping)
    ):
        raise ValueError("Mapping must map known fields to distinct source columns")
    if any(h not in headers for h in mapping.values()):
        raise ValueError("Mapped column missing from file")
    rows = []
    errors = []
    seen = set()
    for index, raw in enumerate(records, 2):
        number = getattr(raw, "number", index)
        data = {
            field: raw[column] for field, column in mapping.items() if raw[column] != ""
        }
        try:
            row = model.model_validate(data).model_dump(mode="json")
            references(kind, row)
            if row["code"] in seen:
                raise ValueError("Duplicate code within file")
            seen.add(row["code"])
            rows.append({"row": number, "data": row})
        except ValidationError as exc:
            for err in exc.errors(include_input=False, include_url=False):
                errors.append(
                    {
                        "row": number,
                        "field": ".".join(map(str, err["loc"])),
                        "code": "invalid_value",
                        "message": err["msg"],
                    }
                )
        except ValueError as exc:
            errors.append(
                {
                    "row": number,
                    "field": "code/references",
                    "code": "duplicate_or_reference",
                    "message": str(exc),
                }
            )
    return rows, errors


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    value = str(value)
    return (
        "'" + value
        if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r"))
        else value
    )


def workbook_bytes(sheets):
    b = Workbook(write_only=True)
    for name, rows in sheets.items():
        ws = b.create_sheet(name)
        for row in rows:
            ws.append([safe_cell(v) for v in row])
    out = io.BytesIO()
    b.save(out)
    return out.getvalue()


def csv_bytes(rows):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerows([[safe_cell(x) for x in r] for r in rows])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def template(kind, format):
    samples = {
        "site": ["SITE-NEW", "示例工厂", "factory", "22.55", "114.06", "CN"],
        "order": [
            "ORDER-NEW",
            "示例订单",
            "SITE-NEW",
            "PRODUCT-01",
            "示例客户",
            "10000.00",
            "CNY",
            "2026-11-01T09:00:00+08:00",
            "pending",
        ],
        "shipment": [
            "SHIP-NEW",
            "示例运输",
            "SITE-NEW",
            "SITE-NEW",
            "ORDER-NEW",
            "road",
            "2026-10-30T09:00:00+08:00",
        ],
        "event": [
            "EVENT-NEW",
            "示例异常",
            "SITE-NEW",
            "ORDER-NEW",
            "high",
            "供给异常，待人工核实",
            "2026-10-17T09:00:00+08:00",
        ],
    }
    rows = [list(ROWS[kind].model_fields), samples[kind]]
    return csv_bytes(rows) if format == "csv" else workbook_bytes({"Import": rows})
