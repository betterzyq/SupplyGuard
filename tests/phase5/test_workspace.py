import hashlib
import io
import json
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
import uuid
import pytest
from openpyxl import load_workbook, Workbook
from sqlalchemy import text
from packages.operations.files import read_table, normalize, template, csv_bytes

BASE = "/api/v1/workspace"


def call(env, method, path, body=None, who="a_enterprise", key=None, **kw):
    headers = {**env.headers[who]}
    if method in {"post", "patch"}:
        headers["Idempotency-Key"] = key or uuid.uuid4().hex
    return env.client.request(method, BASE + path, headers=headers, json=body, **kw)


def preview(
    env, kind, body=None, who="a_enterprise", mapping=None, key=None, format="csv"
):
    content = body if body is not None else template(kind, format)
    headers = read_table(content, "test." + format)[0]
    fields = {x: x for x in headers} if mapping is None else mapping
    return env.client.post(
        BASE + "/imports/preview",
        headers={**env.headers[who], "Idempotency-Key": key or uuid.uuid4().hex},
        data={"kind": kind, "source_type": "synthetic", "mapping": json.dumps(fields)},
        files={"file": ("test." + format, content)},
    )


def imported(env, kind, **kw):
    p = preview(env, kind, **kw)
    assert p.status_code == 200, p.text
    assert p.json()["state"] == "validated", p.text
    r = call(
        env,
        "post",
        "/imports/" + p.json()["id"] + "/confirm",
        who=kw.get("who", "a_enterprise"),
    )
    assert r.status_code == 200, r.text
    return p.json(), r.json()


def demo(env):
    r = call(env, "post", "/demo-scenario", {"scenario_id": "port-congestion"})
    assert r.status_code == 200, r.text
    return r.json()["events"][0]["id"]


def task(env):
    record = demo(env)
    p = {
        "record_id": record,
        "title": "保护交付",
        "due_at": (datetime.now(timezone.utc) + timedelta(days=2)).isoformat(),
        "assignee_id": str(env.users["a_enterprise"]),
        "measures": "核实运输并协调替代供给",
    }
    r = call(env, "post", "/tasks", p)
    assert r.status_code == 200, r.text
    return r.json()


def test_auth_roles_and_tenant_not_client_controlled(enterprise_clients):
    e = enterprise_clients
    assert e.client.get(BASE + "/schema").status_code == 401
    assert call(e, "get", "/schema", who="a_external_supplier").status_code == 403
    assert call(e, "get", "/schema", who="a_admin").status_code == 200
    assert (
        call(
            e,
            "post",
            "/demo-scenario",
            {"scenario_id": "port-congestion", "tenant_id": str(e.tenant_b)},
        ).status_code
        == 422
    )


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_import_idempotent_mapped_and_persistent(enterprise_clients, format):
    e = enterprise_clients
    p, r = imported(e, "site", format=format)
    assert r["inserted"] == 1
    again = call(e, "post", "/imports/" + p["id"] + "/confirm")
    assert again.status_code == 200 and again.json()["replayed"]
    assert call(e, "get", "/records").json()["total"] == 1
    assert call(e, "get", "/records", who="a_admin").json()["total"] == 1
    assert call(e, "get", "/records", who="b_admin").json()["total"] == 0
    # A fresh HTTP client uses the same DB; no browser/session memory is required.
    with e.engine.connect() as c:
        assert (
            c.execute(
                text("select count(*) from sg5_records where tenant_id=:tenant"),
                {"tenant": e.tenant_a},
            ).scalar_one()
            == 1
        )


def test_actual_mapping_and_errors_do_not_commit_partial_rows(enterprise_clients):
    e = enterprise_clients
    malformed = preview(e, "site", mapping={"code": ["not-a-column"]})
    assert (
        malformed.status_code == 422
        and malformed.json()["error"]["code"] == "invalid_mapping"
    )
    data = "站点编号,站点名称,类型,纬度,经度,国家\nS1,工厂,factory,22,114,CN\nS2,错误,factory,91,114,CN\n"
    mapping = dict(
        zip(
            ["code", "name", "site_type", "latitude", "longitude", "country"],
            ["站点编号", "站点名称", "类型", "纬度", "经度", "国家"],
        )
    )
    p = preview(e, "site", body=data.encode(), mapping=mapping)
    assert p.status_code == 200 and p.json()["state"] == "invalid"
    assert any(x["row"] == 3 and x["field"] == "latitude" for x in p.json()["errors"])
    assert call(e, "post", "/imports/" + p.json()["id"] + "/confirm").status_code == 409
    assert call(e, "get", "/records").json()["total"] == 0


def test_duplicate_and_foreign_references_block_all_rows(enterprise_clients):
    e = enterprise_clients
    imported(e, "site", who="b_enterprise")
    p = preview(e, "order")
    assert (
        p.json()["state"] == "invalid"
        and p.json()["errors"][0]["code"] == "reference_not_found"
    )
    assert call(e, "post", "/imports/" + p.json()["id"] + "/confirm").status_code == 409
    header, row = read_table(template("site", "csv"), "test.csv")
    content = csv_bytes([header, *[[r[h] for h in header] for r in row] * 2])
    p = preview(e, "site", body=content)
    assert p.json()["state"] == "invalid"
    assert call(e, "get", "/records").json()["total"] == 0
    assert call(e, "get", "/imports", who="b_admin").json()["total"] == 1


def test_confirmation_revalidates_and_rolls_back_race(enterprise_clients):
    e = enterprise_clients
    p = preview(e, "site").json()
    header, rows = read_table(template("site", "csv"), "test.csv")
    rows[0]["name"] = "先写入的不同名称"
    changed = csv_bytes([header, [rows[0][h] for h in header]])
    imported(e, "site", body=changed)
    response = call(e, "post", "/imports/" + p["id"] + "/confirm")
    assert response.status_code == 409
    assert call(e, "get", "/records").json()["total"] == 1


def test_task_lifecycle_history_cost_and_optimistic_lock(enterprise_clients):
    e = enterprise_clients
    t = task(e)
    for action, extra in [
        ("start", {}),
        (
            "submit",
            {"result": "改道已协调，等待交付验证", "actual_cost_cny": "1140.00"},
        ),
        ("complete", {"result": "人工复核：交付已确认；不代表系统自动执行"}),
    ]:
        r = call(
            e,
            "patch",
            "/tasks/" + t["id"],
            {"expected_version": t["version"], "action": action, **extra},
        )
        assert r.status_code == 200, r.text
        t = r.json()
    assert (
        t["state"] == "completed"
        and t["actual_cost_cny"] == "1140.00"
        and t["version"] == 4
    )
    h = call(e, "get", "/tasks/" + t["id"] + "/history").json()
    assert h["total"] == 4 and h["items"][0]["detail"]["after"]["state"] == "completed"
    stale = call(
        e, "patch", "/tasks/" + t["id"], {"expected_version": 1, "action": "start"}
    )
    assert (
        stale.status_code == 409 and stale.json()["error"]["code"] == "version_conflict"
    )
    assert call(e, "get", "/tasks/" + t["id"], who="b_admin").status_code == 404
    assert (
        call(e, "get", "/tasks/" + t["id"] + "/history", who="b_enterprise").status_code
        == 404
    )


def test_invalid_transitions_assignment_cancel_and_rework(enterprise_clients):
    e = enterprise_clients
    t = task(e)
    assert (
        call(
            e,
            "patch",
            "/tasks/" + t["id"],
            {"expected_version": 1, "action": "complete", "result": "not started"},
        ).status_code
        == 409
    )
    assert (
        call(
            e,
            "patch",
            "/tasks/" + t["id"],
            {
                "expected_version": 1,
                "action": "assign",
                "assignee_id": str(e.users["b_admin"]),
            },
        ).status_code
        == 422
    )
    assert (
        call(
            e, "patch", "/tasks/" + t["id"], {"expected_version": 1, "action": "cancel"}
        ).status_code
        == 409
    )
    r = call(
        e,
        "patch",
        "/tasks/" + t["id"],
        {"expected_version": 1, "action": "cancel", "result": "重复风险，取消任务"},
    )
    assert r.json()["state"] == "cancelled"


def test_same_key_conflict_and_concurrent_versions(enterprise_clients):
    e = enterprise_clients
    t = task(e)
    key = uuid.uuid4().hex
    b = {"expected_version": 1, "action": "start"}
    first = call(e, "patch", "/tasks/" + t["id"], b, key=key)
    second = call(e, "patch", "/tasks/" + t["id"], b, key=key)
    assert first.json() == second.json()
    assert (
        call(
            e,
            "patch",
            "/tasks/" + t["id"],
            {**b, "action": "cancel", "result": "different"},
            key=key,
        ).status_code
        == 409
    )
    t = first.json()

    def update(_):
        return call(
            e,
            "patch",
            "/tasks/" + t["id"],
            {
                "expected_version": t["version"],
                "action": "submit",
                "result": "待验证",
                "actual_cost_cny": "0.00",
            },
        ).status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(update, range(2))) == [200, 409]


def test_risk_export_real_bytes_tenant_hash_expiry_and_formula_safety(
    enterprise_clients,
):
    e = enterprise_clients
    task(e)
    r = call(e, "post", "/reports", {"kind": "risk", "format": "xlsx"})
    assert r.status_code == 200, r.text
    r = r.json()
    # Demo events overlap; each at-risk order is included once.
    assert r["summary"]["exposure_amount_cny"] == "3380000.00"
    data = call(e, "get", "/reports/" + r["id"] + "/download")
    assert (
        data.status_code == 200
        and hashlib.sha256(data.content).hexdigest() == r["sha256"]
    )
    book = load_workbook(io.BytesIO(data.content), read_only=True, data_only=False)
    assert set(book.sheetnames) == {"Summary", "Orders", "Events", "Actions", "Lineage"}
    assert len(list(book["Orders"].rows)) - 1 == r["summary"]["unique_at_risk_orders"]
    action_rows = list(book["Actions"].values)
    assert all(len(row) == len(action_rows[0]) for row in action_rows[1:])
    assert action_rows[1][action_rows[0].index("source_type")] == "synthetic"
    book.close()
    assert (
        call(e, "get", "/reports/" + r["id"] + "/download", who="b_admin").status_code
        == 404
    )
    with e.engine.begin() as c:
        c.execute(
            text(
                "update sg5_reports set expires_at=now()-interval '1 day' where id=:id"
            ),
            {"id": uuid.UUID(r["id"])},
        )
    assert call(e, "get", "/reports/" + r["id"] + "/download").status_code == 410


def test_local_subscription_replay_and_isolation(enterprise_clients):
    e = enterprise_clients
    sub = call(
        e,
        "post",
        "/subscriptions",
        {"name": "本地验证", "topics": ["task.created", "task.updated"]},
    ).json()
    task(e)
    first = call(e, "get", "/subscriptions/" + sub["id"] + "/events").json()
    assert first["external_sent"] is False and len(first["items"]) == 1
    assert first == call(e, "get", "/subscriptions/" + sub["id"] + "/events").json()
    assert (
        call(
            e, "get", f"/subscriptions/{sub['id']}/events?after={first['next_cursor']}"
        ).json()["items"]
        == []
    )
    assert (
        call(
            e, "get", "/subscriptions/" + sub["id"] + "/events", who="b_admin"
        ).status_code
        == 404
    )


def test_file_bounds_formula_and_no_naive_dates():
    with pytest.raises(ValueError):
        read_table(b"xx", "bad.xls")
    with pytest.raises(ValueError, match="Malformed"):
        read_table(b"not a zip", "bad.xlsx")
    book = Workbook()
    book.active.append(["code"])
    book.active.append(["=1+1"])
    b = io.BytesIO()
    book.save(b)
    with pytest.raises(ValueError, match="Formula"):
        read_table(b.getvalue(), "bad.xlsx")
    headers, rows = read_table(template("order", "csv"), "order.csv")
    rows[0]["promised_at"] = "2026-11-01T09:00:00"
    _, errors = normalize("order", headers, rows, {h: h for h in headers})
    assert any(e["field"] == "promised_at" for e in errors)
    assert b"'=SUM" in csv_bytes([["=SUM(1,2)"]])


def test_all_import_entities_references_and_equal_content_dedup(enterprise_clients):
    e = enterprise_clients
    for kind in ["site", "order", "shipment", "event"]:
        p, r = imported(e, kind)
        assert r["inserted"] == 1
    assert call(e, "get", "/records?limit=2&offset=0").json()["total"] == 4
    assert len(call(e, "get", "/records?limit=2&offset=2").json()["items"]) == 2
    assert (
        call(e, "get", "/records?kind=order&q=ORDER-NEW").json()["items"][0]["payload"][
            "amount_cny"
        ]
        == "10000.00"
    )
    # Byte-different file, same normalized record: no silent update, no duplicate.
    content = template("site", "csv") + b"\r\n"
    p, r = imported(e, "site", body=content)
    assert r["inserted"] == 0
    assert call(e, "get", "/records").json()["total"] == 4


def test_pending_verification_rework_and_required_cost(enterprise_clients):
    e = enterprise_clients
    t = task(e)
    t = call(
        e, "patch", "/tasks/" + t["id"], {"expected_version": 1, "action": "start"}
    ).json()
    assert (
        call(
            e,
            "patch",
            "/tasks/" + t["id"],
            {"expected_version": 2, "action": "submit", "result": "缺少费用"},
        ).status_code
        == 422
    )
    t = call(
        e,
        "patch",
        "/tasks/" + t["id"],
        {
            "expected_version": 2,
            "action": "submit",
            "result": "人工完成措施",
            "actual_cost_cny": "0.00",
        },
    ).json()
    assert t["state"] == "pending_verification"
    t = call(
        e,
        "patch",
        "/tasks/" + t["id"],
        {"expected_version": 3, "action": "reopen", "result": "缺少交付凭证，需要补充"},
    ).json()
    assert t["state"] == "in_progress"


def test_private_simulation_explicit_share_and_public_prediction_capture(
    enterprise_clients, monkeypatch, tmp_path
):
    e = enterprise_clients
    monkeypatch.setattr(
        "packages.business.simulation_store.STORE_ROOT", tmp_path / "private"
    )
    simulation = e.client.post(
        "/api/v1/simulation/run",
        headers=e.headers["a_enterprise"],
        json={"preset_id": "hand-check"},
    )
    assert simulation.status_code == 200, simulation.text
    identifier = simulation.json()["run_id"]
    # Same-tenant admin cannot access someone else's private source snapshot.
    assert (
        call(
            e,
            "post",
            "/captures",
            {"kind": "simulation", "identifier": identifier, "share_with_tenant": True},
            who="a_admin",
        ).status_code
        == 404
    )
    assert (
        call(
            e,
            "post",
            "/captures",
            {
                "kind": "simulation",
                "identifier": identifier,
                "share_with_tenant": False,
            },
        ).status_code
        == 422
    )
    shared = call(
        e,
        "post",
        "/captures",
        {"kind": "simulation", "identifier": identifier, "share_with_tenant": True},
    )
    assert shared.status_code == 200 and shared.json()["source_type"] == "synthetic"
    assert (
        call(e, "get", "/records/" + shared.json()["id"], who="a_admin").status_code
        == 200
    )
    from packages.prediction.runtime import cases

    capture = call(
        e,
        "post",
        "/captures",
        {
            "kind": "prediction",
            "identifier": cases()[0]["case_id"],
            "share_with_tenant": True,
        },
    )
    assert (
        capture.status_code == 200
        and capture.json()["source_type"] == "public_historical"
    )
    assert capture.json()["payload"]["domain"] == "dataco-retail-historical"


def test_evaluation_export_and_real_report_tool(enterprise_clients):
    e = enterprise_clients
    r = call(e, "post", "/reports", {"kind": "evaluation", "format": "csv"})
    assert r.status_code == 200, r.text
    data = call(e, "get", "/reports/" + r.json()["id"] + "/download")
    assert (
        data.status_code == 200
        and b"public_historical" in data.content
        and b"average_precision" in data.content
    )
    tool = e.client.post(
        "/api/v1/tools/call",
        headers=e.headers["a_enterprise"],
        json={
            "tool_name": "export_risk_report",
            "tool_input": {"idempotency_key": uuid.uuid4().hex},
        },
    )
    assert tool.status_code == 200, tool.text
    assert tool.json()["status"] == "success" and tool.json()["result"]["download_url"]


def test_database_composite_reference_prevents_foreign_tenant_link(enterprise_clients):
    from sqlalchemy.exc import IntegrityError

    e = enterprise_clients
    imported(e, "site")
    imported(e, "site", who="b_enterprise")
    a = call(e, "get", "/records").json()["items"][0]["id"]
    b = call(e, "get", "/records", who="b_admin").json()["items"][0]["id"]
    with pytest.raises(IntegrityError):
        with e.engine.begin() as c:
            c.execute(
                text(
                    "insert into sg5_links (id,tenant_id,source_id,target_id,relation) values (:id,:tenant,:source,:target,'invalid')"
                ),
                {
                    "id": uuid.uuid4(),
                    "tenant": e.tenant_a,
                    "source": uuid.UUID(a),
                    "target": uuid.UUID(b),
                },
            )


def test_report_and_preview_expiry_and_key_required(enterprise_clients):
    e = enterprise_clients
    assert (
        e.client.post(
            BASE + "/demo-scenario",
            headers=e.headers["a_enterprise"],
            json={"scenario_id": "port-congestion"},
        ).status_code
        == 422
    )
    p = preview(e, "site").json()
    with e.engine.begin() as c:
        c.execute(
            text(
                "update sg5_import_batches set expires_at=now()-interval '1 day' where id=:id"
            ),
            {"id": uuid.UUID(p["id"])},
        )
    assert call(e, "post", "/imports/" + p["id"] + "/confirm").status_code == 410
    assert call(e, "get", "/records").json()["total"] == 0


def test_authenticated_enterprise_openapi_is_filtered_and_resolves(enterprise_clients):
    e = enterprise_clients
    assert e.client.get(BASE + "/openapi").status_code == 401
    r = call(e, "get", "/openapi")
    assert r.status_code == 200, r.text
    assert BASE + "/tasks" in r.json()["paths"]
    assert "/api/v1/admin/audit-logs" not in r.json()["paths"]
    assert r.json()["components"]["securitySchemes"]
    assert any(
        p["name"] == "Idempotency-Key" and p["required"]
        for p in r.json()["paths"][BASE + "/tasks"]["post"]["parameters"]
    )


def test_epoch_like_date_is_not_silently_accepted_as_business_timestamp():
    headers, rows = read_table(template("order", "csv"), "order.csv")
    rows[0]["promised_at"] = "20261101"
    _, errors = normalize("order", headers, rows, {h: h for h in headers})
    assert any(e["field"] == "promised_at" for e in errors)


def test_blank_rows_preserve_original_error_row_numbers():
    content = b"code,name,site_type,latitude,longitude,country\n\nSITE-BAD,broken,factory,91,114,CN\n"
    headers, rows = read_table(content, "blank.csv")
    _, errors = normalize("site", headers, rows, {h: h for h in headers})
    assert errors and all(e["row"] == 3 for e in errors)
